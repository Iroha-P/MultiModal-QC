import argparse
import json
from pathlib import Path

import openpyxl


def parse_row(row: dict, video_path: str) -> dict:
    user_msg = (
        "<video>\n"
        "请检查这段钻井视频操作是否合规。"
        f"井位桩号：{row['well_id']}，设计井深：{row['design_depth']}。"
    )
    if "不合格" not in str(row.get("result", "")):
        assistant_msg = (
            f"该钻井视频检查合格。井位桩号{row['well_id']}，"
            f"完钻深度{row.get('actual_depth', '未知')}米，"
            f"量井深度{row.get('measured_depth', '未知')}米，"
            f"达到设计要求{row['design_depth']}，操作规范。"
        )
    else:
        issues = []
        if row.get("blurry"):
            issues.append("视频录制模糊不清")
        if row.get("no_video"):
            issues.append("存在无关视频内容")
        if row.get("similar"):
            issues.append("与其他视频相似，疑似重复")
        if row.get("no_mark"):
            issues.append("视频无标记信息")
        if row.get("incomplete"):
            issues.append("录制过程不完整或存在分段")
        if row.get("unreported"):
            issues.append("未按要求上报")
        if row.get("depth_fail"):
            issues.append(
                f"井深未达设计要求（设计{row['design_depth']}，"
                f"实际{row.get('actual_depth', '未知')}米）"
            )
        if row.get("helmet"):
            issues.append("操作人员未佩戴安全帽")
        if row.get("uniform"):
            issues.append("操作人员未穿工装")
        if row.get("gloves"):
            issues.append("操作人员未佩戴手套")
        remark = row.get("remark", "")
        if remark and not issues:
            issues.append(remark)
        issue_text = "；".join(issues) if issues else "存在不合格项"
        assistant_msg = f"不合格。发现以下问题：{issue_text}。建议整改后重新检查。"
    return {
        "video": video_path,
        "conversations": [
            {"role": "user", "content": user_msg},
            {"role": "assistant", "content": assistant_msg},
        ],
    }


def read_excel(excel_path: str | Path) -> list[dict]:
    wb = openpyxl.load_workbook(excel_path, read_only=True)
    ws = wb[wb.sheetnames[0]]
    rows_data = []
    for i, row in enumerate(ws.iter_rows(values_only=True)):
        if i < 5:
            continue
        vals = [v for v in row]
        if not vals[0] or not str(vals[0]).strip().isdigit():
            continue
        rows_data.append(
            {
                "well_id": str(vals[1] or ""),
                "design_depth": str(vals[2] or ""),
                "actual_depth": vals[3],
                "measured_depth": vals[5],
                "result": str(vals[9] or ""),
                "no_video": str(vals[10] or ""),
                "similar": str(vals[11] or ""),
                "blurry": str(vals[12] or ""),
                "no_mark": str(vals[13] or ""),
                "incomplete": str(vals[14] or ""),
                "unreported": str(vals[15] or ""),
                "depth_fail": str(vals[17] or ""),
                "remark": str(vals[18] or ""),
                "helmet": str(vals[19] or ""),
                "uniform": str(vals[20] or ""),
                "gloves": str(vals[21] or ""),
            }
        )
    wb.close()
    return rows_data


def convert_drilling(excel_path: str | Path, output_path: str | Path, video_dir: str | Path | None = None):
    rows = read_excel(excel_path)
    video_prefix = Path(video_dir).as_posix().rstrip("/") if video_dir else ""
    samples = []
    for i, row in enumerate(rows):
        video_name = f"video_{i + 1:04d}.mp4"
        video_path = f"{video_prefix}/{video_name}" if video_prefix else video_name
        samples.append(parse_row(row, video_path))

    output = Path(output_path)
    output.parent.mkdir(parents=True, exist_ok=True)
    output.write_text(json.dumps(samples, ensure_ascii=False, indent=2), encoding="utf-8")
    print(f"Converted {len(samples)} drilling records -> {output}")
    return samples


def main():
    parser = argparse.ArgumentParser(description="Convert local private drilling Excel records to QA JSON.")
    parser.add_argument("excel_path", help="Private Excel file path, for example private/drilling/source.xlsx")
    parser.add_argument("output_path", help="Output QA JSON path, for example private/drilling/drilling_qa.json")
    parser.add_argument("--video-dir", default="", help="Optional private video directory prefix.")
    args = parser.parse_args()
    convert_drilling(args.excel_path, args.output_path, args.video_dir or None)


if __name__ == "__main__":
    main()
